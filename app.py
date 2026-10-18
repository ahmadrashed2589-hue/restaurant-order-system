import streamlit as st
import pandas as pd
import os
import plotly.express as px
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter

# ==========================================
# ⚙️ إعداد الصفحة
# ==========================================
st.set_page_config(
    page_title="نظام إدارة المطاعم", 
    layout="wide",
    initial_sidebar_state="expanded",
    page_icon="🍽️"
)

# ==========================================
# 🎨 تنسيق CSS
# ==========================================
st.markdown("""
    <style>
    .stApp { direction: rtl; }
    section[data-testid="stSidebar"] {
        right: 0 !important; left: auto !important;
        border: none !important; box-shadow: -2px 0 5px rgba(0,0,0,0.1);
        overflow: hidden !important; z-index: 999999 !important;
    }
    div[data-testid="stSidebarResizeHandle"] { display: none !important; }
    [data-testid="stSidebarCollapsedControl"] { left: auto !important; right: 10px !important; z-index: 1000000 !important; color: white !important; }
    .main .block-container { max-width: 100% !important; padding-right: 2rem !important; padding-left: 2rem !important; }
    .element-container, .stMarkdown, .stDataFrame, .stTextInput, .stNumberInput, .stSelectbox, .stTextArea, .stButton, .stExpander, .stMultiSelect, .stRadio, .stMetric {
        text-align: right !important; direction: rtl !important;
    }
    div[data-testid="stDataFrame"] div[role="grid"] div[role="row"] div {
        text-align: right !important; direction: rtl !important; justify-content: flex-end !important;
    }
    div[data-testid="metric-container"] {
        background-color: #262730; border: 1px solid #444; padding: 10px; border-radius: 10px; text-align: center;
        box-shadow: 0 2px 5px rgba(0,0,0,0.2); min-height: 100px;
    }
    div[data-testid="stMetricLabel"] { width: 100%; text-align: center !important; color: #d3d3d3; font-size: 13px; }
    div[data-testid="stMetricValue"] { width: 100%; text-align: center !important; color: #ff4b4b !important; font-size: 24px; }
    .stButton button { width: 100%; border-radius: 8px; font-weight: bold; height: 45px; }
    input { text-align: right !important; }
    .delete-btn button { border: 1px solid #ff4b4b !important; color: #ff4b4b !important; padding: 0px 5px; }
    ::placeholder { color: #6c757d !important; opacity: 1 !important; }
    </style>
    """, unsafe_allow_html=True)

# ==========================================
# 📦 إدارة الحالة
# ==========================================
DEFAULT_CONFIG = [
    {"title": "دجاج سحب", "inputs": ["وزن", "سعر"]},
    {"title": "دجاج حب", "inputs": ["وزن", "سعر"]},
    {"title": "بطاطا", "inputs": ["عدد", "سعر"]},
    {"title": "كاتشب", "inputs": ["وزن", "سعر"]},
    {"title": "جبنة", "inputs": ["عدد", "سعر"]},
    {"title": "إضافات", "inputs": ["وزن", "سعر"]},
    {"title": "صدر فيليه برازيلي", "inputs": ["وزن", "سعر"]},
    {"title": "صدر فيليه طازج", "inputs": ["وزن", "سعر"]},
]

if 'form_config' not in st.session_state: st.session_state.form_config = DEFAULT_CONFIG
if 'step' not in st.session_state: st.session_state.step = 0
if 'order_data' not in st.session_state: st.session_state.order_data = {}
if 'expander_toggle' not in st.session_state: st.session_state.expander_toggle = 0

# ==========================================
# 🛠️ دوال النظام
# ==========================================

def get_excel_files():
    return [f for f in os.listdir('.') if f.endswith('.xlsx') and not f.startswith('~$')]

def sync_config_with_excel(filename):
    try:
        # استخدام pandas لأنه يغلق الملف تلقائياً
        df = pd.read_excel(filename, sheet_name="الطلبات", header=None)
        if len(df) < 3: return
        main_headers = df.iloc[1].fillna(method='ffill')
        sub_headers = df.iloc[2].fillna('')
        ignore_cols = ["اليوم", "التاريخ", "رقم الفاتورة", "المجموع الكلي", "الخصم", "المدفوع", "المتبقي", "nan", "None"]
        discovered = {}
        for main, sub in zip(main_headers, sub_headers):
            m_str, s_str = str(main).strip(), str(sub).strip()
            if m_str in ignore_cols or m_str == 'nan': continue
            if m_str not in discovered: discovered[m_str] = []
            if s_str and s_str != 'nan' and s_str not in discovered[m_str]: discovered[m_str].append(s_str)
        
        current_titles = [item['title'] for item in st.session_state.form_config]
        added = False
        for title, inputs in discovered.items():
            if title not in current_titles:
                st.session_state.form_config.append({"title": title, "inputs": inputs})
                added = True
            else:
                for conf in st.session_state.form_config:
                    if conf['title'] == title:
                        for inp in inputs:
                            if inp not in conf['inputs']: conf['inputs'].append(inp); added = True
        if added: st.rerun()
    except: pass

def get_clean_dataframe_for_editing(filename):
    try:
        df = pd.read_excel(filename, sheet_name="الطلبات", header=None)
        if len(df) < 3: return pd.DataFrame()
        main_headers = df.iloc[1].fillna(method='ffill')
        sub_headers = df.iloc[2].fillna('')
        
        clean_columns = []
        for m, s in zip(main_headers, sub_headers):
            m_str, s_str = str(m).strip(), str(s).strip()
            if m_str == s_str or s_str == '' or s_str == 'nan': clean_columns.append(m_str)
            else: clean_columns.append(f"{m_str} - {s_str}")
        
        data_df = df.iloc[3:].copy()
        if len(data_df.columns) > len(clean_columns): data_df = data_df.iloc[:, :len(clean_columns)]
        elif len(data_df.columns) < len(clean_columns):
            for i in range(len(clean_columns) - len(data_df.columns)): data_df[len(data_df.columns)] = 0
        
        data_df.columns = clean_columns
        for col in data_df.columns: data_df[col] = pd.to_numeric(data_df[col], errors='ignore')
        
        # إعادة الحساب (جمع الأسعار فقط)
        price_cols = [c for c in data_df.columns if "سعر" in str(c)]
        for col in price_cols: data_df[col] = pd.to_numeric(data_df[col], errors='coerce').fillna(0)
        calc_total = 0
        if price_cols: calc_total = data_df[price_cols].sum(axis=1)

        if "المجموع الكلي" in data_df.columns: data_df["المجموع الكلي"] = calc_total
        else: data_df.insert(3, "المجموع الكلي", calc_total)
        
        if "الخصم" not in data_df.columns: data_df.insert(4, "الخصم", 0.0)
        else: data_df["الخصم"] = pd.to_numeric(data_df["الخصم"], errors='coerce').fillna(0)
        if "المدفوع" not in data_df.columns: data_df.insert(5, "المدفوع", 0.0)
        else: data_df["المدفوع"] = pd.to_numeric(data_df["المدفوع"], errors='coerce').fillna(0)

        data_df["المتبقي"] = data_df["المجموع الكلي"] - data_df["الخصم"] - data_df["المدفوع"]
        return data_df.iloc[:, ::-1].reset_index(drop=True)
    except: return pd.DataFrame()

def get_financial_transactions(filename):
    wb = None
    try:
        # 🔴 الإصلاح: استخدام load_workbook وإغلاقه يدوياً
        wb = load_workbook(filename, read_only=True, data_only=True)
        if "العمليات_المالية" not in wb.sheetnames: 
            wb.close()
            return pd.DataFrame(columns=["التاريخ", "النوع", "المبلغ", "ملاحظات"])
        
        # قراءة البيانات يدوياً لضمان الدقة وإغلاق الملف
        ws = wb["العمليات_المالية"]
        data = ws.values
        cols = next(data) # الترويسة
        data = list(data)
        # التأكد من الترويسة
        required_cols = ["التاريخ", "النوع", "المبلغ", "ملاحظات"]
        # تحويل لـ DataFrame
        if data:
            df = pd.DataFrame(data, columns=cols)
        else:
            df = pd.DataFrame(columns=required_cols)
            
        # التحقق من وجود الأعمدة المطلوبة، إذا لا نرجع فارغ
        if not all(col in df.columns for col in required_cols):
             wb.close()
             return pd.DataFrame(columns=required_cols)
             
        wb.close() # 🔴 إغلاق الملف ضروري جداً
        return df.iloc[:, ::-1] # قلب الجدول للعرض
    except: 
        if wb: wb.close()
        return pd.DataFrame(columns=["التاريخ", "النوع", "المبلغ", "ملاحظات"])

def save_financial_transaction(filename, date_str, type_, amount, note):
    wb = None
    try:
        if not os.path.exists(filename): return False
        wb = load_workbook(filename)
        if "العمليات_المالية" not in wb.sheetnames:
            ws = wb.create_sheet("العمليات_المالية")
            ws.append(["التاريخ", "النوع", "المبلغ", "ملاحظات"])
            ws.sheet_view.rightToLeft = True
            for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid"); cell.alignment = Alignment(horizontal='center')
        else: ws = wb["العمليات_المالية"]
        ws.append([date_str, type_, amount, note])
        wb.save(filename)
        wb.close() # 🔴 إغلاق
        return True
    except: 
        if wb: wb.close()
        return False

def save_formatted_excel(filename, data_dict, restaurant_name, is_new_file=False, override_data=None):
    all_orders = []
    existing_financials = []
    
    # 1. سحب البيانات المالية القديمة
    if not is_new_file and os.path.exists(filename):
        try:
            # نستخدم دالة get_financial_transactions المعدلة التي تغلق الملف
            fin_df = get_financial_transactions(filename) 
            # ملاحظة: get_financial_transactions ترجع الجدول مقلوباً، نعيده لأصله للحفظ
            if not fin_df.empty: 
                fin_df_original = fin_df.iloc[:, ::-1]
                existing_financials = fin_df_original.values.tolist()
        except: pass

    # 2. تحديد مصدر البيانات
    if override_data is not None:
        all_orders = override_data
    elif not is_new_file and os.path.exists(filename):
        try:
            old_df_view = get_clean_dataframe_for_editing(filename)
            old_df_original = old_df_view.iloc[:, ::-1]
            if not old_df_original.empty: all_orders = old_df_original.to_dict('records')
        except: pass
        if data_dict: all_orders.append(data_dict)
    elif data_dict: all_orders.append(data_dict)
    
    # 3. إعداد الأعمدة
    current_config_map = {item['title']: item['inputs'] for item in st.session_state.form_config}
    for record in all_orders:
        for key in record.keys():
            if " - " in str(key):
                parts = str(key).split(" - ")
                if len(parts) >= 2:
                    title = parts[0]; inp_type = parts[1]
                    if title not in current_config_map: current_config_map[title] = []
                    if inp_type not in current_config_map[title]: current_config_map[title].append(inp_type)

    final_config_for_save = [{"title": k, "inputs": v} for k, v in current_config_map.items()]

    # 4. إعادة الحساب (جمع فقط)
    for record in all_orders:
        row_total = 0.0
        for key, val in record.items():
            if str(key).endswith(" - سعر"):
                try: row_total += float(val)
                except: pass
        
        record["المجموع الكلي"] = row_total
        try: paid = float(record.get("المدفوع", 0))
        except: paid = 0
        try: disc = float(record.get("الخصم", 0))
        except: disc = 0
        
        record["المدفوع"] = paid
        record["الخصم"] = disc
        record["المتبقي"] = row_total - paid - disc

    # 5. الحفظ
    if os.path.exists(filename):
        try: os.remove(filename)
        except PermissionError: return "FILE_OPEN"
            
    wb = Workbook()
    if "Sheet" in wb.sheetnames: del wb["Sheet"]
    ws = wb.create_sheet("الطلبات"); ws.sheet_view.rightToLeft = True
    
    header_fill = PatternFill(start_color="E6B8B7", end_color="E6B8B7", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    title_font = Font(bold=True, size=16); bold_font = Font(bold=True, size=12)
    red_font = Font(color="FF0000", bold=True); green_font = Font(color="008000", bold=True); blue_font = Font(color="0000FF", bold=True)
    
    static_cols = ["اليوم", "التاريخ", "رقم الفاتورة", "المجموع الكلي", "الخصم", "المدفوع", "المتبقي"]
    total_cols = len(static_cols) + sum(len(x['inputs']) for x in final_config_for_save)
    
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title_cell = ws.cell(row=1, column=1, value=restaurant_name)
    title_cell.font = title_font; title_cell.fill = header_fill; title_cell.alignment = center_align
    
    col_idx = 1
    for col in static_cols:
        ws.merge_cells(start_row=2, start_column=col_idx, end_row=3, end_column=col_idx)
        cell = ws.cell(row=2, column=col_idx, value=col)
        cell.fill = header_fill; cell.alignment = center_align; cell.border = thin_border; cell.font = bold_font
        ws.cell(row=3, column=col_idx).border = thin_border
        col_idx += 1
        
    for item in final_config_for_save:
        main_title = item["title"]; inputs = item["inputs"]
        ws.merge_cells(start_row=2, start_column=col_idx, end_row=2, end_column=col_idx + len(inputs) - 1)
        main_cell = ws.cell(row=2, column=col_idx, value=main_title)
        main_cell.fill = header_fill; main_cell.alignment = center_align; main_cell.border = thin_border; main_cell.font = bold_font
        for sub in inputs:
            sub_cell = ws.cell(row=3, column=col_idx, value=sub)
            sub_cell.alignment = center_align; sub_cell.border = thin_border
            sub_cell.font = Font(bold=True, color="FF0000" if "سعر" in sub else "000000")
            col_idx += 1
            
    row_idx = 4
    for record in all_orders:
        c = 1
        ws.cell(row=row_idx, column=c, value=record.get("اليوم")); c+=1
        ws.cell(row=row_idx, column=c, value=record.get("التاريخ")); c+=1
        ws.cell(row=row_idx, column=c, value=record.get("رقم الفاتورة")); c+=1
        ws.cell(row=row_idx, column=c, value=record.get("المجموع الكلي")).font = red_font; c+=1
        ws.cell(row=row_idx, column=c, value=record.get("الخصم")).font = blue_font; c+=1
        ws.cell(row=row_idx, column=c, value=record.get("المدفوع")).font = green_font; c+=1
        ws.cell(row=row_idx, column=c, value=record.get("المتبقي")).font = red_font; c+=1
        for item in final_config_for_save:
            main_title = item["title"]
            for sub in item["inputs"]:
                full_key = f"{main_title} - {sub}"; val = record.get(full_key, 0)
                if pd.isna(val) or val == 'nan': val = 0
                cell = ws.cell(row=row_idx, column=c, value=val)
                if "سعر" in sub: cell.font = red_font
                c += 1
        row_idx += 1

    for row in ws.iter_rows(min_row=4, max_row=max(4, row_idx-1), min_col=1, max_col=total_cols):
        for cell in row: cell.alignment = center_align; cell.border = thin_border
    for i in range(1, total_cols + 1): ws.column_dimensions[get_column_letter(i)].width = 14

    ws_fin = wb.create_sheet("العمليات_المالية"); ws_fin.sheet_view.rightToLeft = True
    ws_fin.append(["التاريخ", "النوع", "المبلغ", "ملاحظات"])
    for cell in ws_fin[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid"); cell.alignment = Alignment(horizontal='center')
    for row in existing_financials: ws_fin.append(row)
        
    try: 
        wb.save(filename)
        wb.close()
        return True
    except PermissionError: return "FILE_OPEN"

def get_all_restaurants_data():
    files = get_excel_files()
    summary_list = []
    for file in files:
        rest_name = file.replace('.xlsx', '')
        clean_df_view = get_clean_dataframe_for_editing(file)
        
        total_sales = clean_df_view["المجموع الكلي"].sum() if not clean_df_view.empty and "المجموع الكلي" in clean_df_view.columns else 0.0
        paid_orders = clean_df_view["المدفوع"].sum() if not clean_df_view.empty and "المدفوع" in clean_df_view.columns else 0.0
        disc_orders = clean_df_view["الخصم"].sum() if not clean_df_view.empty and "الخصم" in clean_df_view.columns else 0.0
        
        fin_df = get_financial_transactions(file)
        paid_ledger = fin_df[fin_df["النوع"] == "دفعة مستلمة"]["المبلغ"].sum() if not fin_df.empty else 0.0
        disc_ledger = fin_df[fin_df["النوع"] == "خصم"]["المبلغ"].sum() if not fin_df.empty else 0.0
        
        total_paid = paid_orders + paid_ledger
        total_discount = disc_orders + disc_ledger
        balance = total_sales - total_paid - total_discount
        summary_list.append({
            "اسم المطعم": rest_name, "إجمالي المبيعات": total_sales,
            "إجمالي الخصومات": total_discount, "إجمالي المقبوضات": total_paid, "الرصيد المتبقي (له)": balance, "الملف": file
        })
    return pd.DataFrame(summary_list)

ARABIC_DAYS = {'Saturday': 'السبت', 'Sunday': 'الأحد', 'Monday': 'الاثنين', 'Tuesday': 'الثلاثاء', 'Wednesday': 'الأربعاء', 'Thursday': 'الخميس', 'Friday': 'الجمعة'}

# ==========================================
# 🎛️ القائمة الجانبية (ثابتة)
# ==========================================
st.sidebar.title("📌 القائمة الرئيسية")

expander_label = f"⚙️ إدارة قائمة الأصناف (عام) {' ' * st.session_state.expander_toggle}"
with st.sidebar.expander(expander_label, expanded=False):
    st.caption("تعديل الأصناف هنا ينطبق على الملفات الجديدة")
    indices_to_remove = []
    for i, item in enumerate(st.session_state.form_config):
        col_txt, col_btn = st.columns([0.85, 0.15])
        with col_txt: st.markdown(f"🔹 **{item['title']}**")
        with col_btn: 
            if st.button("❌", key=f"del_global_{i}"): indices_to_remove.append(i)
    if indices_to_remove:
        for i in sorted(indices_to_remove, reverse=True): del st.session_state.form_config[i]
        st.session_state.expander_toggle += 1; st.rerun()

    st.markdown("---")
    new_item_title = st.text_input("اسم الصنف:", placeholder="مثال: برجر", key="new_global")
    new_item_types = st.multiselect("الخيارات:", ["سعر", "وزن", "عدد"], default=["سعر", "عدد"], key="types_global")
    if st.button("إضافة للقائمة ➕", key="add_global"):
        if new_item_title and new_item_types:
            st.session_state.form_config.append({"title": new_item_title, "inputs": new_item_types})
            st.success(f"تمت إضافة {new_item_title}")
            st.session_state.expander_toggle += 1; st.rerun()

st.sidebar.markdown("---")
app_mode = st.sidebar.radio("تنقل بين الصفحات:", ["🏠 تسجيل الطلبات", "📊 النظرة العامة"])

# ==========================================
# 📊 Dashboard
# ==========================================
if app_mode == "📊 النظرة العامة":
    st.title("📊 لوحة القيادة العامة للمطاعم")
    st.markdown("---")
    if st.button("🔄 تحديث البيانات"): st.rerun()
    all_data_df = get_all_restaurants_data()
    if not all_data_df.empty:
        total_all_sales = all_data_df["إجمالي المبيعات"].sum()
        total_all_discounts = all_data_df["إجمالي الخصومات"].sum()
        total_all_payments = all_data_df["إجمالي المقبوضات"].sum()
        total_all_balance = all_data_df["الرصيد المتبقي (له)"].sum()
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("💰 المبيعات", f"{total_all_sales:,.2f} دينار")
        c2.metric("✂️ الخصومات", f"{total_all_discounts:,.2f} دينار")
        c3.metric("📥 المقبوضات", f"{total_all_payments:,.2f} دينار")
        c4.metric("📉 الديون (لك)", f"{total_all_balance:,.2f} دينار", delta_color="inverse")
        st.markdown("---")
        st.subheader("📈 مقارنة المديونية حسب المطعم")
        fig = px.bar(all_data_df, x="اسم المطعم", y=["إجمالي المبيعات", "الرصيد المتبقي (له)"], barmode='group', color_discrete_sequence=["#3498db", "#e74c3c"])
        fig.update_layout(font=dict(family="Arial", size=14, color="white"), paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)')
        st.plotly_chart(fig, use_container_width=True)
        st.subheader("📋 تفاصيل الحسابات")
        display_df = all_data_df.drop(columns=["الملف"])
        numeric_cols = ["إجمالي المبيعات", "إجمالي الخصومات", "إجمالي المقبوضات", "الرصيد المتبقي (له)"]
        format_dict = {col: "{:,.2f}" for col in numeric_cols}
        st.dataframe(display_df.style.format(format_dict), use_container_width=True)
    else: st.info("لا توجد بيانات.")

# ==========================================
# 🏠 صفحة التسجيل
# ==========================================
elif app_mode == "🏠 تسجيل الطلبات":
    st.sidebar.markdown("---")
    st.sidebar.header("📁 الملفات")
    files = get_excel_files()
    menu_options = files + ["➕ إنشاء ملف جديد"]
    index_to_select = 0
    if 'new_created_file' in st.session_state and st.session_state.new_created_file in menu_options:
        index_to_select = menu_options.index(st.session_state.new_created_file)
    selected_option = st.sidebar.selectbox("اختر المطعم:", menu_options, index=index_to_select)
    restaurant_name_display = ""
    current_file = None

    if selected_option == "➕ إنشاء ملف جديد":
        st.sidebar.markdown("---")
        new_filename = st.sidebar.text_input("اسم الملف (انجليزي):")
        restaurant_header_name = st.sidebar.text_input("اسم المطعم (عربي):")
        if st.sidebar.button("حفظ وإنشاء"):
             if new_filename:
                 fname = new_filename + ".xlsx" if not new_filename.endswith('.xlsx') else new_filename
                 res = save_formatted_excel(fname, None, restaurant_header_name, True)
                 if res == True:
                     st.session_state.new_created_file = fname
                     st.success("تم!"); st.rerun()
                 elif res == "FILE_OPEN": st.error("⚠️ الملف مفتوح! أغلقه أولاً.")
                 else: st.error("خطأ غير معروف")
    else:
        if 'new_created_file' in st.session_state: del st.session_state['new_created_file']
        current_file = selected_option
        try: wb = load_workbook(current_file); ws = wb["الطلبات"]; restaurant_name_display = ws['A1'].value or current_file.replace('.xlsx', '')
        except: restaurant_name_display = current_file.replace('.xlsx', '')

        if current_file:
            sync_config_with_excel(current_file)
            st.markdown(f"<h1 style='text-align: center; color: #ff4b4b;'>{restaurant_name_display}</h1>", unsafe_allow_html=True)
            
            clean_df_view = get_clean_dataframe_for_editing(current_file)
            fin_df = get_financial_transactions(current_file)
            
            ledger_paid = fin_df[fin_df["النوع"] == "دفعة مستلمة"]["المبلغ"].sum() if not fin_df.empty else 0
            ledger_disc = fin_df[fin_df["النوع"] == "خصم"]["المبلغ"].sum() if not fin_df.empty else 0
            
            current_orders_sales = 0; current_orders_paid = 0; current_orders_disc = 0; current_count = 0
            
            if not clean_df_view.empty:
                with st.expander("📝 سجل الطلبات (اضغط للتعديل والدفع)", expanded=True):
                    editor_key = f"ed_{current_file}_{os.path.getmtime(current_file)}"
                    edited_history = st.data_editor(
                        clean_df_view, num_rows="dynamic", use_container_width=True, key=editor_key,
                        column_config={
                            "المجموع الكلي": st.column_config.NumberColumn(disabled=True),
                            "المتبقي": st.column_config.NumberColumn(disabled=True),
                            "المدفوع": st.column_config.NumberColumn(required=True, step=1.0),
                            "الخصم": st.column_config.NumberColumn(required=True, step=0.5)
                        }
                    )
                    
                    # --- 🟢 حساب المجموع (جمع الأسعار فقط) من الجدول المعروض ---
                    price_cols_in_view = [c for c in edited_history.columns if "سعر" in str(c)]
                    if price_cols_in_view:
                        temp_df = edited_history[price_cols_in_view].apply(pd.to_numeric, errors='coerce').fillna(0)
                        current_orders_sales = temp_df.sum(axis=1).sum()
                    else:
                        current_orders_sales = 0
                    
                    current_orders_paid = pd.to_numeric(edited_history["المدفوع"], errors='coerce').fillna(0).sum()
                    current_orders_disc = pd.to_numeric(edited_history["الخصم"], errors='coerce').fillna(0).sum()
                    current_count = len(edited_history)

                    if st.button("💾 حفظ تعديلات السجل"):
                        # ملاحظة: نمرر الداتا المقلوبة (كما في الشاشة) والدالة save_formatted_excel ستتكفل بالباقي
                        res = save_formatted_excel(current_file, None, restaurant_name_display, override_data=edited_history.to_dict('records'))
                        if res == True: st.success("تم التحديث!"); st.rerun()
                        elif res == "FILE_OPEN": st.error("⚠️ الملف مفتوح! يرجى إغلاقه في Excel والمحاولة مرة أخرى.")
                        else: st.error("حدث خطأ أثناء الحفظ.")
            else: st.info("لا توجد طلبات.")

            FINAL_SALES = current_orders_sales
            FINAL_PAID = current_orders_paid + ledger_paid
            FINAL_DISC = current_orders_disc + ledger_disc
            FINAL_REM = FINAL_SALES - FINAL_PAID - FINAL_DISC

            st.markdown("---")
            st.markdown("### 💰 الخلاصة المالية (مباشر)")
            c1, c2, c3, c4, c5 = st.columns(5)
            c1.metric("عدد الفواتير", current_count)
            c2.metric("إجمالي الفواتير", f"{FINAL_SALES:,.2f}")
            c3.metric("مجموع الدفعات", f"{FINAL_PAID:,.2f}")
            c4.metric("مجموع الخصومات", f"{FINAL_DISC:,.2f}")
            c5.metric("الرصيد المتبقي", f"{FINAL_REM:,.2f}", delta_color="inverse")
            st.markdown("---")
            
            with st.expander("➕ إضافة دفعة أو خصم جديد (دفتر الحسابات)", expanded=False):
                ec1, ec2, ec3, ec4 = st.columns([1, 1, 1, 1])
                with ec1: trans_type = st.selectbox("النوع:", ["دفعة مستلمة", "خصم"])
                with ec2: trans_amount = st.number_input("المبلغ:", min_value=0.0, step=1.0)
                with ec3: trans_note = st.text_input("ملاحظة (اختياري):")
                with ec4: 
                    st.write(""); st.write("")
                    if st.button("💾 حفظ العملية"):
                        if trans_amount > 0:
                            save_financial_transaction(current_file, datetime.now().strftime("%Y-%m-%d"), trans_type, trans_amount, trans_note)
                            st.success("تم الحفظ!"); st.rerun()
                        else: st.warning("المبلغ يجب أن يكون أكبر من 0")
            
            if not fin_df.empty:
                with st.expander("📋 سجل العمليات المالية (دفعات منفصلة وخصومات)"):
                    st.dataframe(fin_df.iloc[:, ::-1], use_container_width=True)

            st.markdown("---")
            CURRENT_CONFIG = st.session_state.form_config
            if len(CURRENT_CONFIG) > 0: st.progress((st.session_state.step / (len(CURRENT_CONFIG) + 1)))

            if st.session_state.step >= len(CURRENT_CONFIG):
                st.header("🛒 فاتورة جديدة")
                df_rev = pd.DataFrame(st.session_state.order_data.items(), columns=["البند", "القيمة"])
                edited_df = st.data_editor(df_rev, use_container_width=True, hide_index=True)
                st.session_state.order_data.update(dict(zip(edited_df["البند"], edited_df["القيمة"])))
                
                # --- 🔴 حساب الفاتورة الحالية (جمع فقط) ---
                tot = 0.0
                for item in CURRENT_CONFIG:
                    title = item['title']
                    p_key = f"{title} - سعر"
                    price = st.session_state.order_data.get(p_key)
                    if price: tot += price

                st.info(f"المطلوب: {tot:.2f} دينار")
                
                c1, c2, c3 = st.columns(3)
                inv = c1.text_input("رقم الفاتورة", value=st.session_state.get('inv',''))
                discount_val = c2.number_input("خصم (على الفاتورة)", value=None, step=0.5, placeholder="0.00")
                paid_now = c3.number_input("مدفوع (على الفاتورة)", value=None, step=1.0, placeholder="0.00")
                
                if st.button("✅ حفظ الفاتورة", type="primary"):
                    if not inv: st.error("رقم الفاتورة مطلوب")
                    else:
                        act_paid = paid_now if paid_now is not None else 0.0
                        act_disc = discount_val if discount_val is not None else 0.0
                        fin = {
                            "اليوم": datetime.now().strftime("%A"), "التاريخ": datetime.now().strftime("%Y-%m-%d"),
                            "رقم الفاتورة": inv, "المجموع الكلي": tot, "الخصم": act_disc, "المدفوع": act_paid, 
                            "المتبقي": tot - act_paid - act_disc, **st.session_state.order_data
                        }
                        res = save_formatted_excel(current_file, fin, restaurant_name_display)
                        if res == True:
                            st.balloons(); st.success("تم الحفظ!"); st.session_state.order_data={}; st.session_state.step=0; st.rerun()
                        elif res == "FILE_OPEN": st.error("⚠️ الملف مفتوح! أغلقه أولاً.")
                        else: st.error("حدث خطأ أثناء الحفظ.")
                if st.button("تعديل"): st.session_state.step -= 1; st.rerun()
            else:
                curr = CURRENT_CONFIG[st.session_state.step]
                st.subheader(f"🔽 {curr['title']}")
                with st.form("step_form"):
                    cols = st.columns(2); tmp = {}
                    for i, inp in enumerate(curr["inputs"]):
                        k = f"{curr['title']} - {inp}"; stored = st.session_state.order_data.get(k)
                        disp = stored if (stored and stored > 0) else None
                        val = cols[i%2].number_input(inp, value=disp, step=0.5, key=k, placeholder="0.00")
                        if val is not None: tmp[k] = val
                    st.write("---")
                    c1, c2 = st.columns([1,4])
                    if st.session_state.step > 0: 
                        if c1.form_submit_button("سابق"): st.session_state.step -= 1; st.rerun()
                    if c2.form_submit_button("تالي"):
                        if tmp: st.session_state.order_data.update(tmp); st.session_state.step += 1; st.rerun()
                        else: st.warning("عبي اشي")